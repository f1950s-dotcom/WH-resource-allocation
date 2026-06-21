"""
Excelエクスポートサービス
openpyxlで以下の3シートを作成:
1. 人別シフト表
2. 工程別シフト表
3. コストサマリ
"""
from typing import List, Dict
import io

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from models import OptimizationAssignment, Employee, Process, OptimizationResult


GRAY_FILL = PatternFill(start_color="AAAAAA", end_color="AAAAAA", fill_type="solid")
ORANGE_FILL = PatternFill(start_color="FFB347", end_color="FFB347", fill_type="solid")
HEADER_FILL = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
THIN_BORDER = Border(
    left=Side(style="thin"),
    right=Side(style="thin"),
    top=Side(style="thin"),
    bottom=Side(style="thin"),
)


def _apply_header_style(cell):
    cell.fill = HEADER_FILL
    cell.font = HEADER_FONT
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = THIN_BORDER


def generate_excel(
    date: str,
    result: OptimizationResult,
    assignments: List[OptimizationAssignment],
    employees: List[Employee],
    processes: List[Process],
) -> bytes:
    wb = openpyxl.Workbook()

    emp_map = {e.employee_id: e for e in employees}
    proc_map = {p.process_id: p for p in processes}

    # All time slots sorted
    all_slots = sorted(set(a.time_slot_start for a in assignments))

    # Build lookup: (emp_id, slot) -> assignment
    assign_map: Dict = {}
    for a in assignments:
        assign_map[(a.employee_id, a.time_slot_start)] = a

    # ─── Sheet 1: 人別シフト表 ────────────────────────────────────────────────
    ws1 = wb.active
    ws1.title = "人別シフト表"

    # Header row
    ws1.cell(row=1, column=1, value="氏名")
    _apply_header_style(ws1.cell(row=1, column=1))
    for col_idx, slot in enumerate(all_slots, start=2):
        cell = ws1.cell(row=1, column=col_idx, value=slot)
        _apply_header_style(cell)
        ws1.column_dimensions[get_column_letter(col_idx)].width = 8

    ws1.column_dimensions["A"].width = 15

    # Employee rows
    active_emp_ids = sorted(
        set(a.employee_id for a in assignments),
        key=lambda eid: emp_map.get(eid).name if eid in emp_map else eid,
    )

    for row_idx, emp_id in enumerate(active_emp_ids, start=2):
        emp = emp_map.get(emp_id)
        ws1.cell(row=row_idx, column=1, value=emp.name if emp else emp_id)

        for col_idx, slot in enumerate(all_slots, start=2):
            a = assign_map.get((emp_id, slot))
            cell = ws1.cell(row=row_idx, column=col_idx)
            if a:
                if a.slot_type == "WORK" and a.process_id:
                    proc = proc_map.get(a.process_id)
                    cell.value = proc.process_name if proc else a.process_id
                    if a.is_overtime:
                        cell.border = Border(
                            left=Side(style="medium", color="FF6600"),
                            right=Side(style="medium", color="FF6600"),
                            top=Side(style="medium", color="FF6600"),
                            bottom=Side(style="medium", color="FF6600"),
                        )
                elif a.slot_type in ("LUNCH_BREAK", "LEGAL_BREAK"):
                    cell.value = "休憩"
                    cell.fill = GRAY_FILL
                else:
                    cell.value = ""
            else:
                cell.value = ""
            cell.alignment = Alignment(horizontal="center")

    # ─── Sheet 2: 工程別シフト表 ──────────────────────────────────────────────
    ws2 = wb.create_sheet("工程別シフト表")

    ws2.cell(row=1, column=1, value="工程名")
    _apply_header_style(ws2.cell(row=1, column=1))
    for col_idx, slot in enumerate(all_slots, start=2):
        cell = ws2.cell(row=1, column=col_idx, value=slot)
        _apply_header_style(cell)
        ws2.column_dimensions[get_column_letter(col_idx)].width = 12

    ws2.column_dimensions["A"].width = 20

    # Build process -> slot -> [emp names]
    proc_slot_emps: Dict = {}
    for a in assignments:
        if a.slot_type == "WORK" and a.process_id:
            key = (a.process_id, a.time_slot_start)
            emp = emp_map.get(a.employee_id)
            name = emp.name if emp else a.employee_id
            proc_slot_emps.setdefault(key, []).append(name)

    active_proc_ids = sorted(
        set(a.process_id for a in assignments if a.process_id and a.slot_type == "WORK"),
        key=lambda pid: proc_map.get(pid).display_order if pid in proc_map else 0,
    )

    for row_idx, proc_id in enumerate(active_proc_ids, start=2):
        proc = proc_map.get(proc_id)
        ws2.cell(row=row_idx, column=1, value=proc.process_name if proc else proc_id)

        for col_idx, slot in enumerate(all_slots, start=2):
            names = proc_slot_emps.get((proc_id, slot), [])
            cell = ws2.cell(row=row_idx, column=col_idx, value=", ".join(names) if names else "")
            cell.alignment = Alignment(horizontal="center")

    # ─── Sheet 3: コストサマリ ────────────────────────────────────────────────
    ws3 = wb.create_sheet("コストサマリ")

    headers = ["氏名", "通常時間(h)", "残業時間(h)", "通常費用", "残業費用", "合計費用"]
    for col_idx, h in enumerate(headers, start=1):
        cell = ws3.cell(row=1, column=col_idx, value=h)
        _apply_header_style(cell)
        ws3.column_dimensions[get_column_letter(col_idx)].width = 15

    slot_hours = 15 / 60.0

    for row_idx, emp_id in enumerate(active_emp_ids, start=2):
        emp = emp_map.get(emp_id)
        emp_assigns = [a for a in assignments if a.employee_id == emp_id and a.slot_type == "WORK"]

        normal_slots = [a for a in emp_assigns if not a.is_overtime]
        ot_slots = [a for a in emp_assigns if a.is_overtime]

        normal_hours = len(normal_slots) * slot_hours
        ot_hours = len(ot_slots) * slot_hours
        normal_cost = sum(float(a.slot_cost) for a in normal_slots)
        ot_cost = sum(float(a.slot_cost) for a in ot_slots)
        total_cost = normal_cost + ot_cost

        ws3.cell(row=row_idx, column=1, value=emp.name if emp else emp_id)
        ws3.cell(row=row_idx, column=2, value=round(normal_hours, 2))
        ws3.cell(row=row_idx, column=3, value=round(ot_hours, 2))
        ws3.cell(row=row_idx, column=4, value=round(normal_cost, 2))
        ws3.cell(row=row_idx, column=5, value=round(ot_cost, 2))
        ws3.cell(row=row_idx, column=6, value=round(total_cost, 2))

    # Total row
    total_row = len(active_emp_ids) + 2
    ws3.cell(row=total_row, column=1, value="合計")
    ws3.cell(row=total_row, column=2, value=f"=SUM(B2:B{total_row-1})")
    ws3.cell(row=total_row, column=3, value=f"=SUM(C2:C{total_row-1})")
    ws3.cell(row=total_row, column=4, value=f"=SUM(D2:D{total_row-1})")
    ws3.cell(row=total_row, column=5, value=f"=SUM(E2:E{total_row-1})")
    ws3.cell(row=total_row, column=6, value=f"=SUM(F2:F{total_row-1})")

    # Save to bytes
    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.read()
